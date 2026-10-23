# -*- coding: utf-8 -*-
"""Оформление: протокол (.docx), планы работ (.xlsx), личные сообщения участникам."""
import datetime as dt
from collections import defaultdict
from pathlib import Path
from typing import Dict, List

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from meeting import WD_SHORT, Meeting, fmt_date
from schema import Protocol
from verify import TaskCheck, due_sort_key

NAVY = RGBColor(0x1E, 0x2A, 0x4A)
ORANGE = RGBColor(0xE8, 0x68, 0x18)
GREY = RGBColor(0x6B, 0x72, 0x80)


def due_label(t) -> str:
    if not t.due_date:
        return "не назван"
    d = dt.date.fromisoformat(t.due_date)
    return f"{d:%d.%m.%Y} ({WD_SHORT[d.weekday()]})"


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _fix_widths(t, widths_cm):
    """Фиксированные ширины колонок: Word и LibreOffice берут их из tblGrid."""
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w * 567)))
    for row in t.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)


def _para(doc, text, bold=False, size=11, color=None, italic=False, after=4):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic, r.font.size = bold, italic, Pt(size)
    if color is not None:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(after)
    return p


def protocol_docx(protocol: Protocol, checks: List[TaskCheck], meeting: Meeting, org: str, path: Path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)  # альбом: таблица поручений широкая
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)

    _para(doc, f"{org}   ·   протокол подготовлен ботом, проверен человеком: ____________", size=9, color=GREY)
    _para(doc, protocol.title or "Протокол планёрки", bold=True, size=18, color=NAVY, after=2)
    s = meeting.stats()
    _para(doc, f"{fmt_date(protocol.date)} · председатель: {protocol.chair} · реплик: {s['remarks']}, "
               f"из них голосом: {s['voice']} ({s['voice_minutes']} мин)", size=10, color=GREY, after=8)
    _para(doc, "Присутствовали: " + "; ".join(f"{a.name} ({a.role})" for a in protocol.attendees), size=10)

    _para(doc, "Повестка", bold=True, size=13, color=NAVY)
    for i, a in enumerate(protocol.agenda, 1):
        _para(doc, f"{i}. {a}", after=1)
    if protocol.decisions:
        _para(doc, "Решения", bold=True, size=13, color=NAVY)
        for d in protocol.decisions:
            doc.add_paragraph(d, style="List Bullet")

    _para(doc, "Поручения", bold=True, size=13, color=NAVY)
    flagged = {c.n: c for c in checks}
    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for cell, h in zip(t.rows[0].cells, ["№", "Поручение", "Ответственный", "Срок", "Проверка"]):
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(cell, "1E2A4A")
    for i, task in enumerate(protocol.tasks, 1):
        c = flagged.get(i)
        row = t.add_row().cells
        row[0].text = str(i)
        row[1].text = task.task
        row[2].text = ", ".join(task.assignees) or "—"
        row[3].text = due_label(task)
        row[4].text = "✓" if c is None or c.ok else "⚠ " + "; ".join(c.flags)
        if c is not None and not c.ok:
            _shade(row[4], "FDF0E6")
    _fix_widths(t, [1.0, 12.3, 4.6, 3.0, 5.2])
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)

    if protocol.open_questions:
        _para(doc, "Открытые вопросы (не стали поручениями)", bold=True, size=13, color=NAVY)
        for q in protocol.open_questions:
            doc.add_paragraph(f"{q.question} — {q.why_open}", style="List Bullet")

    bad = [c for c in checks if not c.ok]
    _para(doc, "Проверка ИИ", bold=True, size=13, color=NAVY)
    if bad:
        _para(doc, f"Отмечено поручений: {len(bad)} из {len(checks)}. Проверьте их по стенограмме перед рассылкой.", color=ORANGE)
    else:
        _para(doc, f"Все {len(checks)} поручений прошли проверку: есть ответственный, срок и цитата из стенограммы.")
    if protocol.skipped:
        _para(doc, "Не вошло в протокол: " + "; ".join(protocol.skipped), size=10, color=GREY, italic=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def plans_by_person(protocol: Protocol) -> Dict[str, list]:
    plans = defaultdict(list)
    for i, t in enumerate(protocol.tasks, 1):
        for a in t.assignees or ["(не назначен)"]:
            plans[a].append((i, t))
    for a in plans:
        plans[a].sort(key=lambda it: due_sort_key(it[1]))
    return dict(plans)


def plans_xlsx(protocol: Protocol, checks: List[TaskCheck], path: Path):
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1E2A4A")
    warn_fill = PatternFill("solid", fgColor="FDF0E6")
    thin = Side(style="thin", color="D5D9E3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    flags = {c.n: c.flags for c in checks}

    def sheet(ws, title, rows):
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=13, color="1E2A4A")
        ws.append([])
        ws.append(["№ в протоколе", "Поручение", "Ответственный", "Срок", "Как прозвучало", "Статус", "Проверка ИИ", "Цитата"])
        for cell in ws[3]:
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for n, t in rows:
            d = dt.date.fromisoformat(t.due_date) if t.due_date else None
            f = flags.get(n, [])
            ws.append([n, t.task, ", ".join(t.assignees) or "—", d, t.due_text, "не начато",
                       "⚠ " + "; ".join(f) if f else "✓", t.quote])
            r = ws.max_row
            ws.cell(r, 4).number_format = "DD.MM.YYYY"
            if f:
                ws.cell(r, 7).fill = warn_fill
        for row in ws.iter_rows(min_row=3):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col, w in zip("ABCDEFGH", [8, 55, 24, 12, 18, 12, 30, 45]):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A4"

    ws = wb.active
    ws.title = "Все поручения"
    sheet(ws, f"{protocol.title} — все поручения", sorted(enumerate(protocol.tasks, 1), key=lambda it: due_sort_key(it[1])))
    for person, items in plans_by_person(protocol).items():
        name = person.replace("/", " ")[:31]
        sheet(wb.create_sheet(name), f"План работ: {person}", items)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def person_message(person: str, items: list, protocol: Protocol) -> str:
    lines = [f"📋 {person}, ваши поручения по итогам планёрки {fmt_date(protocol.date)}:", ""]
    for n, t in items:
        others = [a for a in t.assignees if a != person]
        co = f" (вместе с {', '.join(others)})" if others else ""
        lines.append(f"• до {due_label(t)} — {t.task}{co}")
    lines += ["", "Если срок или формулировка неверны — ответьте председателю до конца дня."]
    return "\n".join(lines)


def summary_message(protocol: Protocol, checks: List[TaskCheck]) -> str:
    bad = [c for c in checks if not c.ok]
    lines = [f"✅ Протокол готов: {protocol.title}, {fmt_date(protocol.date)}",
             f"Поручений: {len(protocol.tasks)}, решений: {len(protocol.decisions)}, открытых вопросов: {len(protocol.open_questions)}."]
    if bad:
        lines.append(f"⚠ Требуют проверки: {len(bad)}")
        for c in bad:
            t = protocol.tasks[c.n - 1]
            lines.append(f"  {c.n}. {t.task[:70]} — {'; '.join(c.flags)}")
    for q in protocol.open_questions:
        lines.append(f"❓ {q.question} — {q.why_open}")
    return "\n".join(lines)
