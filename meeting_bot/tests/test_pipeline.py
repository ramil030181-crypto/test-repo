# -*- coding: utf-8 -*-
"""Проверка конвейера без интернета: LLM_BACKEND=fixture, распознавание подменено."""
import asyncio
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import openpyxl
import pytest
from docx import Document

HERE = Path(__file__).resolve().parents[1]
TRANSCRIPT = HERE.parent / "Входные_данные" / "Планёрка_23.11.2026_стенограмма.txt"
FIXTURE = HERE / "tests" / "fixture_protocol_2026-11-23.json"


@pytest.fixture(autouse=True)
def offline_env(monkeypatch):
    monkeypatch.setenv("LLM_BACKEND", "fixture")
    monkeypatch.setenv("LLM_FIXTURE", str(FIXTURE))


@pytest.fixture
def directory():
    from meeting import Directory
    return Directory(HERE / "participants.json")


def test_speaker_from_first_word(directory):
    assert directory.split_speaker("Ахметов. Пальцы придут завтра") == ("Ахметов Р.С.", "Пальцы придут завтра")
    assert directory.split_speaker("Менеджер ТО: чек-лист закрыт") == ("Менеджер по ТО", "чек-лист закрыт")
    assert directory.split_speaker("Байжанов Ерлан Кайратович, начинаем")[0] == "Байжанов Е.К."
    # обращение по имени — не говорящий
    assert directory.split_speaker("Ирина, что по бюджету?")[0] is None
    assert directory.split_speaker("Так, коллеги, начинаем")[0] is None
    # Whisper потерял точку после фамилии
    assert directory.split_speaker("Сейткали заявка готова") == ("Сейткали Д.М.", "заявка готова")
    assert directory.split_speaker("Менеджер должен подготовить")[0] is None


def test_calendar_hint_weekdays():
    from meeting import calendar_hint
    hint = calendar_hint("2026-11-23")
    assert "23.11.2026, понедельник" in hint
    assert "пт 27.11" in hint and "вт 01.12" in hint


def test_offline_run_on_reference(tmp_path):
    import offline
    res = offline.main([str(TRANSCRIPT), "--date", "2026-11-23", "--data", str(tmp_path)])
    assert len(res.protocol.tasks) == 16
    flagged = {c.n: c.flags for c in res.checks if not c.ok}
    # без срока — ровно 2 поручения (заказ после визы, PT-1104); выдуманных цитат нет
    assert set(flagged) == {3, 10}
    assert all("срок не назван" in f[0] for f in flagged.values())
    assert res.docx.exists() and res.xlsx.exists()

    text = "\n".join(p.text for p in Document(res.docx).paragraphs)
    assert "Открытые вопросы" in text and "E-301C" in text
    cells = [c.text for t in Document(res.docx).tables for r in t.rows for c in r.cells]
    assert "23.11.2026 (пн)" in cells

    wb = openpyxl.load_workbook(res.xlsx)
    assert "Менеджер по ТО" in wb.sheetnames and "Все поручения" in wb.sheetnames
    me = wb["Менеджер по ТО"]
    dates = [r[3].value for r in me.iter_rows(min_row=4) if r[3].value]
    assert len(dates) == 7 and dates == sorted(dates)  # личный план — по срокам

    msg = res.personal["Литвинова И."]
    assert "27.11.2026" in msg and "вместе с Менеджер по ТО" in msg
    # стенограмма: 26 реплик, шапка не попала
    assert res.transcript.read_text(encoding="utf-8").count("\n") == 26


def test_checks_catch_model_mistakes(directory):
    from meeting import Meeting
    from schema import Protocol
    from verify import check

    m = Meeting("2026-11-23")
    for line in TRANSCRIPT.read_text(encoding="utf-8").splitlines()[1:]:
        who, text = directory.split_speaker(line)
        m.add(who, text)
    p = Protocol.model_validate_json(FIXTURE.read_text(encoding="utf-8"))
    p.tasks[0].quote = "Менеджер ТО, закажите новый насос до понедельника"  # выдумано
    p.tasks[1].due_date = "2026-11-20"  # раньше планёрки
    p.tasks[3].due_date = "2026-11-28"  # суббота
    p.tasks[4].assignees = ["Петров А.А."]  # нет такого
    p.tasks[5].assignees = []
    flags = {c.n: " ".join(c.flags) for c in check(p, m, directory)}
    assert "выдумано" in flags[1]
    assert "раньше даты планёрки" in flags[2]
    assert "выходной" in flags[4]
    assert "нет в справочнике" in flags[5]
    assert "нет ответственного" in flags[6]
    assert flags[7] == ""


def test_service_add_undo_bind(tmp_path, directory):
    from service import MeetingService
    s = MeetingService(tmp_path, directory)
    s.start(1, "23.11.2026 Планёрка")
    _, who = s.add(1, 42, "Иванов Пётр", "Всем привет")
    assert who == "Иванов Пётр"
    assert s.bind(42, "Ахметов") == "Ахметов Р.С."
    _, who = s.add(1, 42, "Иванов Пётр", "Обогрев сделаем")
    assert who == "Ахметов Р.С."  # после /iam
    _, who = s.add(1, 42, "Иванов Пётр", "Ким. Двигатель греется")
    assert who == "Ким В.А."  # фамилия в начале реплики важнее привязки
    assert s.undo(1).startswith("Ким В.А.")
    assert len(s.active(1).remarks) == 2
    assert json.loads((tmp_path / "bindings.json").read_text(encoding="utf-8")) == {"Ахметов Р.С.": 42}


def test_schema_is_strict_enough_for_structured_output():
    from schema import Protocol
    sch = Protocol.model_json_schema()
    assert set(sch["required"]) >= {"tasks", "open_questions", "date"}
    task = sch["$defs"]["Task"]
    assert set(task["required"]) == {"task", "assignees", "due_date", "due_text", "quote"}


# --- Telegram-слой на подставных объектах (без сети) ------------------------

def _update(chat=100, uid=7, text=None, voice=None, first="Атлас", last=None):
    msg = MagicMock()
    msg.text, msg.chat_id = text, chat
    msg.voice, msg.audio, msg.video_note, msg.document = voice, None, None, None
    msg.reply_text = AsyncMock()
    u = SimpleNamespace(effective_chat=SimpleNamespace(id=chat), message=msg, effective_message=msg,
                        effective_user=SimpleNamespace(id=uid, first_name=first, last_name=last, username=None))
    return u


def test_bot_flow_text_voice_protocol(tmp_path, monkeypatch):
    import bot
    from service import MeetingService
    monkeypatch.setattr(bot, "SERVICE", MeetingService(tmp_path, bot.DIRECTORY))
    monkeypatch.setattr(bot, "ALLOWED", {7})
    monkeypatch.setattr(bot.asr, "transcribe", lambda path, size, hint: ("Ахметов. Пальцы муфты придут завтра", 4.2))

    ctx = SimpleNamespace(args=["23.11.2026"], bot=MagicMock())
    ctx.bot.send_chat_action = AsyncMock()
    ctx.bot.send_document = AsyncMock()
    ctx.bot.send_message = AsyncMock()

    async def run():
        await bot.cmd_new(_update(text="/new"), ctx)
        for line in TRANSCRIPT.read_text(encoding="utf-8").splitlines()[1:]:
            await bot.on_text(_update(text=line), ctx)
        voice = MagicMock()
        voice.get_file = AsyncMock(return_value=MagicMock(download_to_drive=AsyncMock()))
        up = _update(voice=voice)
        await bot.on_audio(up, ctx)
        assert "Ахметов Р.С. (4 с)" in up.message.reply_text.call_args.args[0]

        stranger = _update(uid=999, text="привет")
        await bot.on_text(stranger, ctx)
        assert "Нет доступа" in stranger.message.reply_text.call_args.args[0]

        await bot.cmd_iam(_update(text="/iam", uid=7), SimpleNamespace(args=["Литвинова"], bot=ctx.bot))
        up = _update(text="/protocol")
        await bot.cmd_protocol(up, ctx)
        assert ctx.bot.send_document.await_count == 3  # протокол, планы, стенограмма
        q = SimpleNamespace(data="send", answer=AsyncMock(), from_user=SimpleNamespace(id=7),
                            message=SimpleNamespace(chat_id=100, reply_text=AsyncMock()))
        await bot.on_button(SimpleNamespace(callback_query=q), ctx)
        uid, text = ctx.bot.send_message.call_args.args
        assert uid == 7 and text.startswith("📋 Литвинова И.")

    asyncio.run(run())
    m = bot.SERVICE.active(100)
    assert len(m.remarks) == 27 and m.remarks[-1].source == "voice"


def test_build_app_registers_handlers():
    import bot
    app = bot.build_app("123456:TEST-token-not-real")
    assert sum(len(h) for h in app.handlers.values()) == 11
